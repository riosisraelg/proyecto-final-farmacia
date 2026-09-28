#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera la Presentacion del Proyecto Final en PowerPoint (.pptx).

Estilo: MINIMALISTA MONOCROMATICO — texto NEGRO sobre fondo BLANCO.
Unico elemento institucional: el LOGO de Tecmilenio SOLO en la portada.

OUTLINE (15 diapositivas) definido por el alumno, alineado al documento
"Presentacion del Proyecto Final" y a la rubrica (100 pts). Mismo contenido
y orden que docs/latex/slides/slides.tex.

Uso:    .venv/bin/python docs/gen_pptx.py
Salida: docs/latex/slides/slides.pptx
"""

import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# --- Rutas -----------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
ATT = os.path.join(ROOT, "attachments")
LOGO = os.path.join(ATT, "Tecmilenio Logo.png")
ER = os.path.join(ATT, "WhatsApp Image 2026-09-25 at 09.53.27.jpeg")
OUT = os.path.join(HERE, "latex", "slides", "slides.pptx")

# --- Paleta monocromatica -------------------------------------------
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x40, 0x40, 0x40)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
MARGIN = Inches(0.7)
CONTENT_W = SW - 2 * MARGIN

_num = [0]


def add_slide():
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return s


def add_title(slide, text):
    box = slide.shapes.add_textbox(MARGIN, Inches(0.4), CONTENT_W, Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = BLACK
    line = slide.shapes.add_connector(1, MARGIN, Inches(1.28),
                                      MARGIN + Inches(5.2), Inches(1.28))
    line.line.color.rgb = BLACK
    line.line.width = Pt(1.5)
    return box


def _add_runs(paragraph, text):
    for tok in re.split(r"(\*\*.*?\*\*|`.*?`)", text):
        if not tok:
            continue
        run = paragraph.add_run()
        if tok.startswith("**") and tok.endswith("**"):
            run.text = tok[2:-2]
            run.font.bold = True
            run.font.name = FONT
        elif tok.startswith("`") and tok.endswith("`"):
            run.text = tok[1:-1]
            run.font.name = "Courier New"
        else:
            run.text = tok
            run.font.name = FONT
        run.font.color.rgb = BLACK


def add_body(slide, blocks, top=Inches(1.5), left=MARGIN, width=CONTENT_W,
             size=18, height=Inches(5.4), sub_size=None):
    """blocks: lista de dict {'label': str|None, 'items': [str,...]} o str simple.
    label -> subtitulo en negrita; items -> vinetas con guion."""
    sub_size = sub_size or size
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for blk in blocks:
        if isinstance(blk, str):
            blk = {"items": [blk]}
        label = blk.get("label")
        if label:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(6)
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = label
            r.font.bold = True
            r.font.size = Pt(sub_size)
            r.font.name = FONT
            r.font.color.rgb = BLACK
        for it in blk.get("items", []):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(4)
            p.level = 1 if label else 0
            _add_runs(p, "– " + it)
            for run in p.runs:
                run.font.size = Pt(size)
    return box


def add_note(slide, text, top, size=14):
    box = slide.shapes.add_textbox(MARGIN, top, CONTENT_W, Inches(0.5))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.italic = True
    r.font.size = Pt(size)
    r.font.name = FONT
    r.font.color.rgb = GRAY


def add_evidence_box(slide, caption, top=Inches(5.0), height=Inches(1.7)):
    w = Inches(9.0)
    left = (SW - w) // 2
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = BLACK
    shp.line.width = Pt(1)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"[Evidencia: {caption}]"
    r.font.italic = True
    r.font.size = Pt(15)
    r.font.name = FONT
    r.font.color.rgb = GRAY


def add_footer(slide):
    _num[0] += 1
    box = slide.shapes.add_textbox(MARGIN, SH - Inches(0.5), CONTENT_W, Inches(0.35))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Farmacia con Consultorio Médico"
    r.font.size = Pt(10)
    r.font.name = FONT
    r.font.color.rgb = GRAY
    box2 = slide.shapes.add_textbox(SW - MARGIN - Inches(1.0), SH - Inches(0.5),
                                    Inches(1.0), Inches(0.35))
    p2 = box2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = str(_num[0])
    r2.font.size = Pt(10)
    r2.font.name = FONT
    r2.font.color.rgb = GRAY


# =====================================================================
# 1. PORTADA (unico lugar con el logo)
# =====================================================================
s = add_slide()
with Image.open(LOGO) as im:
    lw, lh = im.size
logo_h = Inches(0.75)
logo_w = Emu(int(logo_h * lw / lh))
s.shapes.add_picture(LOGO, SW - MARGIN - logo_w, Inches(0.5), height=logo_h)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.4), SW - Inches(2.0), Inches(2.8))
tf = tb.text_frame
tf.word_wrap = True
for i, (text, size, bold, color) in enumerate([
    ("Farmacia con Consultorio Médico", 38, True, BLACK),
    ("Gestión de Consultas y Medicamentos Recetados", 20, False, GRAY),
    ("", 12, False, BLACK),
    ("[Integrantes]", 16, False, BLACK),
    ("LSTI2311: Bases de Datos", 16, False, BLACK),
    ("Prof. [Profesor]", 16, False, BLACK),
    ("26 de septiembre de 2026", 13, False, GRAY),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(5)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color

# =====================================================================
# 2. QUE EMPRESA REPRESENTAMOS
# =====================================================================
s = add_slide()
add_title(s, "¿Qué empresa representamos?")
add_body(s, [
    {"label": "Quiénes somos", "items": ["Consultorio médico con farmacia asociada."]},
    {"label": "Qué áreas intervienen",
     "items": ["Consultorio (médicos), farmacia (dispensación), compras (proveedores) y administración."]},
    {"label": "Cómo se relacionan el consultorio y la farmacia",
     "items": ["El médico consulta y emite la receta; la farmacia la surte y registra la venta."]},
])
add_footer(s)

# =====================================================================
# 3. PROBLEMATICA ACTUAL
# =====================================================================
s = add_slide()
add_title(s, "¿Cuál es la problemática actual?")
add_body(s, [
    {"label": "Situación actual", "items": [
        "Desconexión entre la farmacia y el consultorio.",
        "Operan como dos negocios distintos: información y procesos separados.",
        "Experiencia de usuario fragmentada: sin datos de fidelidad ni conocimiento de la conversión de recetas.",
    ]},
    {"label": "Manejo de la información", "items": [
        "La información se administra de forma arcaica.",
        "Uso de papel; datos dispersos, no centralizados ni conectados para generar información estratégica.",
    ]},
    {"label": "Principales problemas identificados", "items": [
        "Ausencia de automatización.",
        "Desabasto de medicamentos clave y mermas por caducidad no detectada.",
        "Procesos manuales que retrasan la reposición y provocan errores de dispensación.",
    ]},
], size=15, sub_size=16, top=Inches(1.4))
add_footer(s)

# =====================================================================
# 4. IMPACTO
# =====================================================================
s = add_slide()
add_title(s, "¿Cuál es el impacto?")
add_body(s, [
    {"label": "Impacto en los pacientes", "items": [
        "Experiencias distintas entre consultorio y farmacia.",
        "Creación de doble perfil.",
    ]},
    {"label": "Impacto en la atención", "items": [
        "Atención lenta y fragmentada.",
    ]},
    {"label": "Impacto en la farmacia", "items": [
        "Desconocimiento del rendimiento del negocio.",
        "Pérdida de ventas.",
        "Ausencia de competitividad.",
    ]},
])
add_footer(s)

# =====================================================================
# 5. CAUSA RAIZ
# =====================================================================
s = add_slide()
add_title(s, "¿Cuál es la causa raíz?")
add_body(s, [
    {"label": "Falta de integración de la información", "items": [
        "Consultorio y farmacia no comparten una misma fuente de datos.",
        "Información dispersa y en papel: se duplica y no se conecta.",
    ]},
    {"label": "Procesos manuales", "items": [
        "Operación dependiente de tareas manuales, sin conexión entre áreas.",
        "No hay un flujo estandarizado de la información.",
    ]},
    {"label": "Falta de validaciones y controles", "items": [
        "Sin reglas de negocio aplicadas de forma automática.",
        "Sin control de inventario, caducidades ni trazabilidad receta–venta.",
    ]},
])
add_footer(s)

# =====================================================================
# 6. ESTRATEGIA
# =====================================================================
s = add_slide()
add_title(s, "¿Qué estrategia debemos seguir?")
add_body(s, [
    "**Entender las reglas del negocio:** qué se valida, qué se controla.",
    "**Mapear el proceso actual:** de la consulta a la venta y la reposición.",
    "**Estandarizar la operación:** un flujo único y ordenado.",
    "**Centralizar la información:** una sola fuente de verdad.",
    "**Automatizar los procesos:** validaciones y registros sin intervención manual.",
], top=Inches(1.6), size=19, height=Inches(4.0))
add_note(s, "Primero la operación, luego la tecnología.", Inches(5.8))
add_footer(s)

# =====================================================================
# 7. SOLUCION TECNOLOGICA
# =====================================================================
s = add_slide()
add_title(s, "¿Cuál es nuestra solución tecnológica?")
add_body(s, [
    "**Base de datos relacional en MySQL:** centraliza toda la información.",
    "**Modelo Entidad-Relación:** 20 tablas con claves y relaciones definidas.",
    "**Integración de consultorio y farmacia:** un solo sistema, sin islas de datos.",
], top=Inches(1.6), size=19, height=Inches(3.4))
add_note(s, "Evidencia: Modelo Entidad-Relación (siguiente diapositiva).", Inches(5.2))
add_footer(s)

# ER a PANTALLA COMPLETA (evidencia, sin titulo ni pie)
s = add_slide()
with Image.open(ER) as im:
    iw, ih = im.size
scale = min(SW / iw, SH / ih)
img_w = Emu(int(iw * scale))
img_h = Emu(int(ih * scale))
s.shapes.add_picture(ER, (SW - img_w) // 2, (SH - img_h) // 2, width=img_w, height=img_h)

# =====================================================================
# 8. COMO SE GESTIONAN LOS DATOS
# =====================================================================
s = add_slide()
add_title(s, "¿Cómo se gestionan los datos?")
add_body(s, [
    {"label": "Flujo de la información", "items": [
        "Paciente → consulta → receta → surtido → venta → inventario → reposición.",
    ]},
    {"label": "Información clínica", "items": [
        "Historial, diagnósticos, alergias y contraindicaciones del paciente.",
    ]},
    {"label": "Recetas y medicamentos", "items": [
        "Recetas vinculadas a consulta, medicamento, dosis y duración.",
    ]},
    {"label": "Inventario y ventas", "items": [
        "Stock por lote, caducidades, ventas y su vínculo con la receta.",
    ]},
], size=16, sub_size=17)
add_footer(s)

# =====================================================================
# 9. COMO AUTOMATIZAMOS LOS PROCESOS
# =====================================================================
s = add_slide()
add_title(s, "¿Cómo automatizamos los procesos?")
add_body(s, [
    "**Vistas:** historial clínico, inventario con alertas y medicamentos más recetados/surtidos.",
    "**Stored Procedures:** registrar consulta/receta, surtir receta y generar reposición.",
    "**Triggers:** validan recetas y disparan la auditoría de inventario.",
    "**LOG de auditoría:** tabla Auditoria_Inventario con cada movimiento de stock.",
], top=Inches(1.5), size=17, height=Inches(3.2))
add_evidence_box(s, "resultados reales de la BD (vistas / SP / triggers)", top=Inches(5.0))
add_footer(s)

# =====================================================================
# 10. COMO APLICAMOS LAS REGLAS DE NEGOCIO
# =====================================================================
s = add_slide()
add_title(s, "¿Cómo aplicamos las reglas de negocio?")
add_body(s, [
    "**Validación de pacientes y médicos:** paciente activo y médico con cédula vigente.",
    "**Alergias y contraindicaciones:** se verifican contra los medicamentos recetados.",
    "**Interacciones medicamentosas:** se bloquean combinaciones peligrosas.",
    "**Validación de recetas:** dosis, duración y vigencia dentro de rango.",
], top=Inches(1.5), size=17, height=Inches(3.2))
add_evidence_box(s, "mensaje real de validación/error", top=Inches(5.0))
add_footer(s)

# =====================================================================
# 11. COMO CONTROLAMOS EL INVENTARIO
# =====================================================================
s = add_slide()
add_title(s, "¿Cómo controlamos el inventario?")
add_body(s, [
    "**Stock de medicamentos:** estado suficiente / bajo / crítico / agotado.",
    "**Lotes y caducidades:** se dispensa el lote no vencido más próximo a caducar.",
    "**Reposición:** órdenes automáticas al llegar al stock mínimo.",
    "**Auditoría de movimientos:** cada cambio de stock queda registrado.",
], top=Inches(1.5), size=17, height=Inches(3.2))
add_evidence_box(s, "inventario + LOG de auditoría", top=Inches(5.0))
add_footer(s)

# =====================================================================
# 12. COMO INTEGRAMOS OPERACION Y TECNOLOGIA
# =====================================================================
s = add_slide()
add_title(s, "¿Cómo integramos operación y tecnología?")
add_body(s, [
    "**Proceso operativo:** el flujo real del negocio, estandarizado.",
    "**Reglas de negocio:** llevadas a restricciones, SP y triggers.",
    "**Base de datos:** MySQL como fuente única de verdad.",
    "**Conexión Java–MySQL:** acceso a los datos desde la aplicación (JDBC).",
], top=Inches(1.5), size=17, height=Inches(3.2))
add_evidence_box(s, "conexión y consulta desde Java (App.java)", top=Inches(5.0))
add_footer(s)

# =====================================================================
# 13. IMPORTANCIA DE UNA ESTRATEGIA INTEGRAL
# =====================================================================
s = add_slide()
add_title(s, "¿Por qué es importante una estrategia integral?")
add_body(s, [
    "**Entender las reglas de negocio:** define qué debe validar el sistema.",
    "**Mapear correctamente el flujo de información:** evita islas de datos.",
    "**Integrar operación y tecnología:** una sola fuente de verdad.",
    "**Gestionar la correcta ejecución:** vistas, SP y triggers hacen cumplir las reglas.",
], top=Inches(1.6), size=19, height=Inches(3.6))
add_note(s, "Responde directamente a los cuatro puntos de la exposición oral.", Inches(5.4))
add_footer(s)

# =====================================================================
# 14. BENEFICIOS
# =====================================================================
s = add_slide()
add_title(s, "¿Qué beneficios obtenemos?")
add_body(s, [
    "**Seguridad del paciente:** validaciones clínicas automáticas.",
    "**Eficiencia operativa:** procesos estandarizados y más rápidos.",
    "**Control del inventario:** menos desabasto y menos mermas.",
    "**Automatización:** menos errores manuales.",
    "**Información para la toma de decisiones:** reportes de consumo y conversión.",
], top=Inches(1.6), size=19, height=Inches(4.6))
add_footer(s)

# =====================================================================
# 15. CONCLUSION
# =====================================================================
s = add_slide()
add_title(s, "Conclusión")
add_body(s, [
    "**Problema identificado:** consultorio y farmacia desconectados, información arcaica.",
    "**Estrategia aplicada:** estandarizar, centralizar y automatizar la operación.",
    "**Solución implementada:** base de datos relacional con vistas, SP y triggers.",
    "**Resultado obtenido:** atención más segura, menos pérdidas y trazabilidad total.",
], top=Inches(1.6), size=19, height=Inches(4.6))
add_footer(s)

# =====================================================================
# GRACIAS
# =====================================================================
s = add_slide()
box = s.shapes.add_textbox(Inches(1.0), Inches(3.0), SW - Inches(2.0), Inches(1.6))
tf = box.text_frame
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Gracias"
r.font.size = Pt(44); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = BLACK
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(10)
r = p.add_run(); r.text = "¿Preguntas?"
r.font.size = Pt(20); r.font.name = FONT; r.font.color.rgb = GRAY

prs.save(OUT)
print(f"OK -> {OUT}  ({len(prs.slides._sldIdLst)} slides)")
